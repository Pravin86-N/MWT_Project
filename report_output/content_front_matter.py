import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def build_front_matter(doc, img_dir):
    logo_path = os.path.join(img_dir, "image1.png")

    # ==================== COVER PAGE ====================
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.paragraph_format.line_spacing = 1.2
    r_title = p_title.add_run("SMART FUEL DISPATCH & FLEET TELEMETRY SYSTEM")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0, 0, 0)

    # Course
    p_course = doc.add_paragraph()
    p_course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_course.paragraph_format.space_before = Pt(6)
    p_course.paragraph_format.space_after = Pt(24)
    r_course = p_course.add_run("23CS54C – MODERN WEB TECHNOLOGIES")
    r_course.font.name = 'Times New Roman'
    r_course.font.size = Pt(13)
    r_course.font.bold = True

    # Submitted by
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(12)
    p_sub.paragraph_format.space_after = Pt(4)
    r_sub = p_sub.add_run("Submitted by")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True

    # Student Name and Reg
    p_student = doc.add_paragraph()
    p_student.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_student.paragraph_format.space_before = Pt(0)
    p_student.paragraph_format.space_after = Pt(28)
    r_student = p_student.add_run("PRAVIN N (24104080)")
    r_student.font.name = 'Times New Roman'
    r_student.font.size = Pt(14)
    r_student.font.bold = True

    # Partial fulfillment text
    p_ful = doc.add_paragraph()
    p_ful.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ful.paragraph_format.space_before = Pt(6)
    p_ful.paragraph_format.space_after = Pt(2)
    r_ful = p_ful.add_run("In partial fulfillment for the award of the degree\nof")
    r_ful.font.name = 'Times New Roman'
    r_ful.font.size = Pt(11)

    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deg.paragraph_format.space_before = Pt(2)
    p_deg.paragraph_format.space_after = Pt(2)
    r_deg = p_deg.add_run("BACHELOR OF TECHNOLOGY\nin\nINFORMATION TECHNOLOGY")
    r_deg.font.name = 'Times New Roman'
    r_deg.font.size = Pt(13)
    r_deg.font.bold = True

    # College Logo
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(20)
        p_logo.paragraph_format.space_after = Pt(16)
        r_logo = p_logo.add_run()
        r_logo.add_picture(logo_path, width=Inches(1.2))

    # College Name
    p_col = doc.add_paragraph()
    p_col.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_col.paragraph_format.space_before = Pt(4)
    p_col.paragraph_format.space_after = Pt(2)
    r_col = p_col.add_run("NATIONAL ENGINEERING COLLEGE")
    r_col.font.name = 'Times New Roman'
    r_col.font.size = Pt(14)
    r_col.font.bold = True

    p_aut = doc.add_paragraph()
    p_aut.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aut.paragraph_format.space_before = Pt(0)
    p_aut.paragraph_format.space_after = Pt(2)
    r_aut = p_aut.add_run("(An Autonomous Institution affiliated to Anna University, Chennai)")
    r_aut.font.name = 'Times New Roman'
    r_aut.font.size = Pt(10.5)

    p_addr = doc.add_paragraph()
    p_addr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_addr.paragraph_format.space_before = Pt(0)
    p_addr.paragraph_format.space_after = Pt(14)
    r_addr = p_addr.add_run("K.R.NAGAR, KOVILPATTI - 628503")
    r_addr.font.name = 'Times New Roman'
    r_addr.font.size = Pt(11)
    r_addr.font.bold = True

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_before = Pt(0)
    p_date.paragraph_format.space_after = Pt(0)
    r_date = p_date.add_run("OCTOBER - 2026")
    r_date.font.name = 'Times New Roman'
    r_date.font.size = Pt(12)
    r_date.font.bold = True

    # ==================== BONAFIDE CERTIFICATE ====================
    doc.add_page_break()

    if os.path.exists(logo_path):
        p_logo2 = doc.add_paragraph()
        p_logo2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo2.paragraph_format.space_before = Pt(10)
        p_logo2.paragraph_format.space_after = Pt(8)
        r_logo2 = p_logo2.add_run()
        r_logo2.add_picture(logo_path, width=Inches(1.0))

    p_col2 = doc.add_paragraph()
    p_col2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_col2.paragraph_format.space_before = Pt(0)
    p_col2.paragraph_format.space_after = Pt(2)
    r_col2 = p_col2.add_run("NATIONAL ENGINEERING COLLEGE")
    r_col2.font.name = 'Times New Roman'
    r_col2.font.size = Pt(13)
    r_col2.font.bold = True

    p_aut2 = doc.add_paragraph()
    p_aut2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aut2.paragraph_format.space_before = Pt(0)
    p_aut2.paragraph_format.space_after = Pt(2)
    r_aut2 = p_aut2.add_run("(An Autonomous Institution affiliated to Anna University, Chennai)\nK.R.NAGAR, KOVILPATTI - 628503")
    r_aut2.font.name = 'Times New Roman'
    r_aut2.font.size = Pt(10)

    p_bon = doc.add_paragraph()
    p_bon.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bon.paragraph_format.space_before = Pt(20)
    p_bon.paragraph_format.space_after = Pt(18)
    r_bon = p_bon.add_run("BONAFIDE CERTIFICATE")
    r_bon.font.name = 'Times New Roman'
    r_bon.font.size = Pt(14)
    r_bon.font.bold = True

    p_cert = doc.add_paragraph()
    p_cert.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_cert.paragraph_format.line_spacing = 1.3
    p_cert.paragraph_format.space_before = Pt(6)
    p_cert.paragraph_format.space_after = Pt(36)
    r_cert = p_cert.add_run(
        "This is to certify that this project report, \"Smart Fuel Dispatch & Fleet Telemetry System\", "
        "is the bonafide work of PRAVIN N (24104080) who carried out the project work under my supervision."
    )
    r_cert.font.name = 'Times New Roman'
    r_cert.font.size = Pt(12)

    # Signature Guide
    p_guide = doc.add_paragraph()
    p_guide.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_guide.paragraph_format.space_before = Pt(20)
    p_guide.paragraph_format.space_after = Pt(36)
    r_guide = p_guide.add_run("Course Instructor / Guide        ")
    r_guide.font.name = 'Times New Roman'
    r_guide.font.size = Pt(11)
    r_guide.font.bold = True

    p_viva = doc.add_paragraph()
    p_viva.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_viva.paragraph_format.space_before = Pt(20)
    p_viva.paragraph_format.space_after = Pt(40)
    r_viva = p_viva.add_run("Submitted to the 23CS54C – MODERN WEB TECHNOLOGIES Viva-Voce examination held at National Engineering College, K.R.Nagar, Kovilpatti on __________________")
    r_viva.font.name = 'Times New Roman'
    r_viva.font.size = Pt(11)

    p_exam = doc.add_paragraph()
    p_exam.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_exam.paragraph_format.space_before = Pt(10)
    p_exam.paragraph_format.space_after = Pt(0)
    r_ex1 = p_exam.add_run("Internal Examiner")
    r_ex1.font.name = 'Times New Roman'
    r_ex1.font.size = Pt(11)
    r_ex1.font.bold = True
    r_sp = p_exam.add_run("\t\t\t\t\t\t\t")
    r_ex2 = p_exam.add_run("Co-Examiner / External Examiner")
    r_ex2.font.name = 'Times New Roman'
    r_ex2.font.size = Pt(11)
    r_ex2.font.bold = True

    # ==================== TABLE OF CONTENTS ====================
    doc.add_page_break()

    p_toc_t = doc.add_paragraph()
    p_toc_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc_t.paragraph_format.space_before = Pt(12)
    p_toc_t.paragraph_format.space_after = Pt(16)
    r_toc_t = p_toc_t.add_run("TABLE OF CONTENTS")
    r_toc_t.font.name = 'Times New Roman'
    r_toc_t.font.size = Pt(15)
    r_toc_t.font.bold = True

    # Table of contents table
    toc_data = [
        ("1", "INTRODUCTION", "3"),
        ("2", "OBJECTIVES", "4"),
        ("3", "DESCRIPTION", "5"),
        ("4", "CLASS DIAGRAM", "7"),
        ("5", "TABLE STRUCTURES", "9"),
        ("6", "ER DIAGRAM", "14"),
        ("7", "TECH STACK", "16"),
        ("8", "MODULES", "19"),
        ("9", "CONCLUSION", "53"),
        ("-", "REFERENCES", "55")
    ]

    toc_table = doc.add_table(rows=len(toc_data) + 1, cols=3)
    toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    toc_table.autofit = False

    # Style TOC Table header
    headers = ["CH NO", "TITLE", "PAGE NO"]
    col_widths = [1.0, 4.27, 1.0]

    header_row = toc_table.rows[0]
    header_tr = header_row._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    header_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    for idx, name in enumerate(headers):
        cell = header_row.cells[idx]
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="EAF0F6"/>'))
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(name)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.bold = True

    for r_idx, row_item in enumerate(toc_data):
        row = toc_table.rows[r_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        for c_idx, val in enumerate(row_item):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)
            if r_idx in [0, 7, 8]:
                r.font.bold = True

    for row in toc_table.rows:
        for idx, w in enumerate(col_widths):
            row.cells[idx].width = Inches(w)

    print("Front matter successfully generated.")

if __name__ == "__main__":
    print("content_front_matter module ready.")
