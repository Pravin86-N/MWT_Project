import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(BASE_DIR, "user_images")
OUT_DOCX = os.path.join(BASE_DIR, "Smart_Fuel_Dispatch_Fleet_Telemetry_Final_Report.docx")

# Import chapter modules
import content_front_matter
import content_chapters_1_to_4
import content_chapters_5_to_7
import content_chapter_8
import content_chapter_9_ref

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="B8C4D0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def add_chapter_header(doc, ch_num, ch_title):
    doc.add_page_break()
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(14)
    p1.paragraph_format.space_after = Pt(4)
    p1.paragraph_format.keep_with_next = True
    r1 = p1.add_run(f"CHAPTER {ch_num}")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(16)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0, 0, 0)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(18)
    p2.paragraph_format.keep_with_next = True
    r2 = p2.add_run(ch_title)
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(16)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0, 0, 0)
    return p2

def add_heading_2(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_subheading(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_bullet(doc, title, text):
    p = doc.add_paragraph(style='List Paragraph')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    if title:
        r1 = p.add_run(title + ": ")
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(12)
        r1.font.bold = True
    r2 = p.add_run(text)
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(12)
    return p

def add_image_with_caption(doc, img_path, caption_text, width_inches=5.8):
    if not img_path or not os.path.exists(img_path):
        return None
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.keep_with_next = True
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(width_inches))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(12)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(10.5)
    r_cap.font.bold = True
    r_cap.font.italic = True
    return p_img

def add_styled_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    header_row = table.rows[0]
    header_tr = header_row._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    header_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    for col_idx, text in enumerate(headers):
        cell = header_row.cells[col_idx]
        set_cell_background(cell, "EAF0F6")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.0)
        r.font.bold = True

    for row_idx, row_data in enumerate(data):
        row = table.rows[row_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            cell = row.cells[col_idx]
            if bg_color != "FFFFFF":
                set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=65, bottom=65, left=100, right=100)
            p = cell.paragraphs[0]
            val_str = str(cell_value)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx != 0 or len(val_str) > 15 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(val_str)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.5)

    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)
    return table

helpers = {
    'add_chapter_header': add_chapter_header,
    'add_heading_2': add_heading_2,
    'add_subheading': add_subheading,
    'add_paragraph': add_paragraph,
    'add_bullet': add_bullet,
    'add_image_with_caption': add_image_with_caption,
    'add_styled_table': add_styled_table
}

def generate_full_report():
    print("Beginning report generation...")
    doc = docx.Document()

    # Section 0: Front Matter
    sec0 = doc.sections[0]
    sec0.page_width = Inches(8.27)
    sec0.page_height = Inches(11.69)
    sec0.top_margin = Inches(1.0)
    sec0.bottom_margin = Inches(1.0)
    sec0.left_margin = Inches(1.0)
    sec0.right_margin = Inches(1.0)
    sec0.header_distance = Inches(0.5)
    sec0.footer_distance = Inches(0.5)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # 1. Build Front Matter
    print("Building Front Matter...")
    content_front_matter.build_front_matter(doc, IMG_DIR)

    # Add section break for Body
    sec1 = doc.add_section()
    sec1.page_width = Inches(8.27)
    sec1.page_height = Inches(11.69)
    sec1.top_margin = Inches(1.0)
    sec1.bottom_margin = Inches(1.0)
    sec1.left_margin = Inches(1.0)
    sec1.right_margin = Inches(1.0)
    sec1.header_distance = Inches(0.5)
    sec1.footer_distance = Inches(0.5)

    # Configure right-aligned PAGE number in footer
    footer = sec1.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run = footer_p.add_run()
    f_run.font.name = 'Times New Roman'
    f_run.font.size = Pt(11)
    
    # Insert dynamic PAGE field in footer
    f_p_elem = footer_p._p
    f_field_xml = parse_xml(r'''
        <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="22"/></w:rPr>
            <w:fldChar w:fldCharType="begin"/>
            <w:instrText xml:space="preserve"> PAGE </w:instrText>
            <w:fldChar w:fldCharType="separate"/>
            <w:fldChar w:fldCharType="end"/>
        </w:r>
    ''')
    f_p_elem.append(f_field_xml)

    # 2. Build Chapters 1 to 4
    print("Building Chapters 1 to 4...")
    content_chapters_1_to_4.build_chapters_1_to_4(doc, IMG_DIR, helpers)

    # 3. Build Chapters 5 to 7
    print("Building Chapters 5 to 7...")
    content_chapters_5_to_7.build_chapters_5_to_7(doc, IMG_DIR, helpers)

    # 4. Build Chapter 8
    print("Building Chapter 8 (Modules)...")
    content_chapter_8.build_chapter_8(doc, IMG_DIR, helpers)

    # 5. Build Chapter 9 and References
    print("Building Chapter 9 and References...")
    content_chapter_9_ref.build_chapter_9_and_references(doc, helpers)

    # Save document
    print(f"Saving final report to {OUT_DOCX}...")
    doc.save(OUT_DOCX)
    print("Document successfully created and saved!")

if __name__ == "__main__":
    generate_full_report()
